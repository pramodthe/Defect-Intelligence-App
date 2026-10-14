import base64
import io
import json
from contextlib import closing

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from qdrant_client import QdrantClient

from app import visual_embeddings as visual
from app.embedding_index import queue_reindex, rebuild
from app.main import create_app
from app.retrieval import evidence_media
from app.store import digest
from tests.conftest import publish_fixture


def shape(color, circle=False):
    image = Image.new('RGB', (256,256), 'white')
    draw = ImageDraw.Draw(image)
    getattr(draw, 'ellipse' if circle else 'rectangle')((40,40,215,215), fill=color)
    stream = io.BytesIO()
    image.save(stream, 'PNG')
    return stream.getvalue()


def report():
    doc = Document()
    doc.add_heading('Visual fixtures', 1)
    for image in [shape('red'), shape('blue', True)]:
        doc.add_paragraph('Site photograph')
        doc.add_picture(io.BytesIO(image))
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def configure(store, provider='fireworks_voyage', enabled=True, dimensions=256):
    store.config.visual_embedding_provider = provider
    store.config.visual_embedding_deployment = 'accounts/synthetic/deployments/voyage-mm'
    store.config.visual_embedding_enabled = enabled
    store.config.visual_embedding_dimensions = dimensions
    store.config.fireworks_api_key = 'synthetic-fireworks-key'
    store.config.voyage_api_key = 'synthetic-voyage-key'


def client(store, requests, responder=None):
    def transport(request):
        payload = json.loads(request.content)
        requests.append((str(request.url), dict(request.headers), payload))
        if responder:
            return responder(payload)
        inputs = payload.get('inputs', payload.get('input'))
        output = []
        for index, entry in enumerate(inputs):
            parts = entry['content']
            image = next((p for p in parts if p['type'] != 'text'), None)
            if image:
                data_url = image.get('image_base64') or image['image_url']['url']
                raw = base64.b64decode(data_url.split(',',1)[1])
                with Image.open(io.BytesIO(raw)) as decoded:
                    red, _, blue = decoded.convert('RGB').getpixel((128,128))
                blue_selected = blue > red
            else:
                blue_selected = 'blue' in parts[0]['text']
            vector = [0.] * store.config.visual_embedding_dimensions
            vector[1 if blue_selected else 0] = 2.
            output.append({'index': index, 'embedding': vector, 'object': 'embedding'})
        return httpx.Response(200, json={'object':'list', 'data':list(reversed(output))})
    return httpx.Client(transport=httpx.MockTransport(transport))


@pytest.mark.parametrize('provider', ['fireworks_voyage', 'voyage'])
def test_synthetic_capability_probe_checks_images_and_queries_without_report_assets(store, provider):
    configure(store, provider, enabled=False)
    requests = []
    with client(store, requests) as transport:
        result = visual.test_connection(store, transport)
    assert result['status'].startswith('Synthetic image/text test passed')
    assert visual.capable(store) and not visual.configured(store)
    assert len(requests) == 2
    assert not store.all('SELECT * FROM assets')
    first, second = requests[0][2], requests[1][2]
    assert first['input_type'] == 'document' and second['input_type'] == 'query'
    if provider == 'fireworks_voyage':
        assert requests[0][0] == 'https://api.fireworks.ai/inference/v1/embeddings'
        assert first['model'] == 'accounts/synthetic/deployments/voyage-mm'
        assert first['input'][0]['content'][0]['image_url']['url'].startswith('data:image/png;base64,')
        assert first['dimensions'] == 256
    else:
        assert requests[0][0] == 'https://api.voyageai.com/v1/multimodalembeddings'
        assert first['inputs'][0]['content'][0]['type'] == 'image_base64'
        assert first['truncation'] is False and first['output_dimension'] == 256
    assert all('image_url' not in json.dumps(p) or 'http' not in json.dumps(p) for _, _, p in requests)


def test_source_pixels_cannot_be_sent_until_enabled_keyed_and_capability_verified(store):
    configure(store, enabled=False)
    requests = []
    with client(store, requests) as transport:
        embedding = visual.VoyageVisualEmbedding(store, client=transport)
        with pytest.raises(ValueError, match='Enable image embeddings'):
            embedding.encode([visual.content(image=shape('red'))])
        store.config.visual_embedding_enabled = True
        with pytest.raises(ValueError, match='Enable image embeddings'):
            visual.VoyageVisualEmbedding(store, client=transport).encode([visual.content(image=shape('red'))])
        store.config.fireworks_api_key = ''
        with pytest.raises(ValueError, match='API key'):
            visual.VoyageVisualEmbedding(store, client=transport).encode([visual.content(image=shape('red'))])
    assert requests == [] and not store.all('SELECT * FROM operations')


def test_text_only_or_identical_response_fails_image_capability_probe(store):
    configure(store)
    requests = []
    def identical(payload):
        inputs = payload.get('inputs', payload.get('input'))
        return httpx.Response(200, json={'data':[{'index':i,'embedding':[1.] + [0.]*255} for i in range(len(inputs))]})
    with client(store, requests, identical) as transport:
        with pytest.raises(ValueError, match='Synthetic image test failed'):
            visual.test_connection(store, transport)
    assert not visual.capable(store) and not visual.configured(store)


def test_timeout_unknown_attempt_is_charged_and_never_replayed(store):
    configure(store)
    requests = []
    with client(store, requests) as transport:
        visual.test_connection(store, transport)
    with client(store, requests, lambda payload: httpx.Response(429, json={'error':'synthetic secret response'})) as transport:
        embedding = visual.VoyageVisualEmbedding(store, client=transport)
        inputs = [visual.content('different synthetic source context', shape('red'))]
        with pytest.raises(RuntimeError, match='no automatic retry') as failure:
            embedding.encode(inputs)
        assert 'synthetic secret response' not in str(failure.value)
        with pytest.raises(RuntimeError, match='unknown outcome'):
            embedding.encode(inputs)
    assert len(requests) == 3
    assert store.require("SELECT COUNT(*) n FROM operations WHERE state='reserved'")['n'] == 1


def test_visual_cache_is_separate_by_input_type_and_remote_urls_are_rejected(store):
    configure(store)
    requests = []
    with client(store, requests) as transport:
        visual.test_connection(store, transport)
        embedding = visual.VoyageVisualEmbedding(store, client=transport)
        inputs = [visual.content('cache text')]
        assert embedding.encode(inputs, 'document') == embedding.encode(inputs, 'document')
        embedding.encode(inputs, 'query')
        assert len(requests) == 4
        with pytest.raises(ValueError, match='remote image URLs'):
            embedding.encode([{'content':[{'type':'image_url','image_url':'https://untrusted.example/file.png'}]}])
        assert len(requests) == 4


def test_visual_settings_persist_without_exposing_keys_and_require_valid_deployment(store):
    configure(store, enabled=False)
    api = TestClient(create_app(store.config))
    payload = {'expected_version':1,'visual_embedding_provider':'fireworks_voyage',
        'visual_embedding_deployment':'accounts/synthetic/deployments/voyage-mm',
        'visual_embedding_dimensions':1024,'visual_embedding_enabled':True}
    response = api.put('/api/visual-embedding-settings', json=payload, headers={'X-Workspace-Request':'1'})
    assert response.status_code == 200
    assert 'synthetic-fireworks-key' not in response.text and 'synthetic-voyage-key' not in response.text
    assert response.json()['configured'] is False
    assert api.put('/api/visual-embedding-settings', json=payload, headers={'X-Workspace-Request':'1'}).status_code == 400
    fresh = TestClient(create_app(store.config)).get('/api/visual-embedding-settings').json()
    assert fresh['visual_embedding_enabled'] is True and fresh['version'] == 2
    with pytest.raises(ValueError, match='deployment path'):
        visual.spec(store.config.model_copy(update={'visual_embedding_deployment':'https://unknown.example'}))


def test_missing_deployment_and_unverified_visual_index_cannot_enqueue_rebuild(store, ontology):
    publish_fixture(store, ontology, report())
    configure(store)
    store.config.qdrant_url = 'synthetic'
    store.config.visual_embedding_deployment = ''
    with pytest.raises(ValueError, match='synthetic image test'):
        queue_reindex(store)
    assert not store.all("SELECT * FROM outbox WHERE kind='embedding_reindex'")


def test_original_images_are_immutable_and_pixel_derivative_is_bounded(store):
    original = shape('red')
    asset = store.asset(original, 'image/png')
    processed = visual.pixels(original)
    assert store.asset_path(asset).read_bytes() == original
    with Image.open(io.BytesIO(processed)) as image:
        assert image.mode == 'RGB' and image.width <= 1600
    with pytest.raises(ValueError, match='image_pixel_limit'):
        visual.pixels(original, max_pixels=1)
    with pytest.raises(ValueError):
        visual.pixels(b'not an image')


def test_visual_projection_rebuild_and_pixel_retrieval_preserve_source_links(store, ontology, monkeypatch):
    raw = report()
    document = publish_fixture(store, ontology, raw)
    original_manifest = store.active_manifest()
    configure(store)
    store.config.qdrant_url = 'isolated-local-qdrant'
    requests = []
    real_embedding = visual.VoyageVisualEmbedding
    with client(store, requests) as transport:
        visual.test_connection(store, transport)
        monkeypatch.setattr(visual, 'VoyageVisualEmbedding', lambda s, owner=None: real_embedding(s, owner, transport))
        monkeypatch.setattr(visual.VisualProjection, 'client', lambda self: QdrantClient(path=str(store.root/'visual-qdrant')))
        queued = queue_reindex(store)
        result = rebuild(store, ontology, store.require('SELECT * FROM outbox WHERE id=?', (queued['id'],)))
        assert result['manifest_id'] != original_manifest
        state = visual.public_settings(store)
        assert state['index_state'] == 'ready' and state['indexed_media'] == 2
        hits = evidence_media(store, ontology, 'blue circle')
        assert hits['items'][0]['asset_id'] == 'asset_' + digest(shape('blue', True))
        assert hits['items'][0]['retrieval_methods'] == ['pixel_semantic']
        assert hits['items'][0]['original_url'].endswith('/content')
        assert 'Multimodal image retrieval' in hits['coverage']
        source_asset = store.require('SELECT asset_id FROM documents WHERE id=?', (document['document_id'],))['asset_id']
        assert store.asset_path(source_asset).read_bytes() == raw
        with closing(QdrantClient(path=str(store.root/'visual-qdrant'))) as qdrant:
            collection = qdrant.get_collections().collections[0].name
            assert qdrant.get_collection(collection).config.params.vectors.size == 256
        blue_id = hits['items'][0]['id']
        store.execute('UPDATE media SET caption=? WHERE id=?', ('Changed source caption', blue_id))
        stale = evidence_media(store, ontology, 'blue circle')
        assert all(i['id'] != blue_id for i in stale['items'])
        assert any('Stale visual vector' in warning for warning in stale['warnings'])
        configure(store, provider='voyage')
        no_cross_profile = evidence_media(store, ontology, 'blue circle')
        assert not any('pixel_semantic' in i['retrieval_methods'] for i in no_cross_profile['items'])


def test_failed_visual_projection_cannot_replace_manifest(store, ontology, monkeypatch):
    publish_fixture(store, ontology, report())
    before = store.active_manifest()
    configure(store)
    store.config.qdrant_url = 'test'
    requests = []
    with client(store, requests) as transport:
        visual.test_connection(store, transport)
    queued = queue_reindex(store)
    def fail(*args, **kwargs):
        raise RuntimeError('synthetic unavailable visual index')
    monkeypatch.setattr(visual.VisualProjection, 'write', fail)
    with pytest.raises(RuntimeError):
        rebuild(store, ontology, store.require('SELECT * FROM outbox WHERE id=?', (queued['id'],)))
    assert store.active_manifest() == before
