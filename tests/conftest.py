import io
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont
from app.settings import Settings
from app.store import Store
from app.ontology import Ontology
from app.ingestion import register, parse_docx, persist_ir
from app.extraction import build_units, annotate_domain
from app.chunking import build_chunks
from app.projections import stage


@pytest.fixture
def store(tmp_path):
    config = Settings(_env_file=None, workspace_root=tmp_path, checkpoint_backend="sqlite-dev",
                      tokenizer_path="", text_provider="disabled", neo4j_uri="", qdrant_url="")
    s = Store(config)
    s.initialize()
    return s


@pytest.fixture
def ontology(store):
    return Ontology(store)


def synthetic_docx(long=False, count=3):
    doc = Document()
    doc.styles.add_style("CHeading", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_heading("External defects", 1)
    doc.add_heading("Planter Box Membrane", 2)
    doc.add_paragraph("Sample Findings:", "CHeading")
    for index in range(count):
        doc.add_paragraph(f"DS-{index+1}, Unit {index+1} planter")
        doc.add_paragraph("Membrane does not extend above the soil line. 😀 e\u0301")
    doc.add_paragraph("Reasoning:", "CHeading")
    doc.add_paragraph("AS 4654.2-2012 Figure 2.17 is cited. The report contains a reference; applicability is unresolved.")
    doc.add_paragraph(("Drainage and membrane termination require inspection. " * (400 if long else 1)))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Measurement"
    table.cell(0, 1).text = "Condition / note"
    table.cell(1, 0).text = "12 mm"
    table.cell(1, 1).text = "Only where the source states this condition"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested exception"
    doc.add_paragraph("Remediation:", "CHeading")
    doc.add_paragraph("Option 1: Enclose corridor ends")
    doc.add_paragraph("Obtain approval before enclosing. Verify ventilation.")
    doc.add_paragraph("Option 2: Waterproof corridor")
    doc.add_paragraph("Obtain regulated design; do not combine with Option 1.")
    image = Image.new("RGB", (900, 200), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 32)
    except OSError:
        font = ImageFont.load_default(size=30)
    draw.text((25, 25), "AS 4654.2-2012", fill="black", font=font)
    draw.text((25, 90), "FIGURE 2.17 TYPICAL PLANTER BOX", fill="black", font=font)
    b = io.BytesIO()
    image.save(b, "PNG")
    b.seek(0)
    doc.add_picture(b)
    doc.add_heading("Concrete finish quality", 2)
    doc.add_paragraph("Reasoning:", "CHeading")
    doc.add_paragraph("AS 3610.1-2010 concrete finish Class 2 is distinct from Building Class 2.")
    doc.add_paragraph("NCC 2016 and NCC 2022 are cited for different discussions.")
    b.seek(0)
    doc.add_picture(b)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def publish_fixture(store, ontology, raw=None):
    r = register(raw or synthetic_docx(), "synthetic.docx", "company_report", store)
    d = store.one("SELECT * FROM documents WHERE id=?", (r["document_id"],))
    persist_ir(store, parse_docx(store, d))
    build_units(store, d["id"], ontology)
    annotate_domain(store, d["id"], ontology)
    build_chunks(store, d["id"], ontology.version)
    projection = stage(store, d["id"], ontology)
    store.activate(projection["projection_id"], {"profile": projection["profile"]})
    store.execute("UPDATE documents SET status='partial' WHERE id=?", (d["id"],))
    return r
