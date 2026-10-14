import json
import io

from fastapi.testclient import TestClient
from docx import Document

from app.main import create_app
from app.report_sections import original_report_sections
from app.retrieval import search, source
from app.store import canonical
from app.ingestion import register, parse_docx, persist_ir
from app.extraction import build_units, annotate_domain
from app.chunking import build_chunks
from app.projections import stage
from tests.conftest import publish_fixture
from tests.test_workflows_api import run_record


def saved_run(store, ontology):
    publish_fixture(store, ontology)
    rid = run_record(store)
    results = search(store, ontology, 'planter membrane reasoning repair')
    # Only one small hit is saved. The output must restore its source section,
    # including authored reasoning and separate options, rather than just its chunk.
    hit = next(e for e in results['items'] if e['metadata']['defect_item_id'])
    store.execute('UPDATE runs SET evidence=? WHERE id=?', (canonical({'items': [hit]}), rid))
    return rid


def test_original_wording_sections_restore_context_and_keep_option_scope(store, ontology):
    rid = saved_run(store, ontology)
    before = store.require('SELECT evidence,answer FROM runs WHERE id=?', (rid,))
    result = original_report_sections(store, rid)
    assert result['wording'] == 'original_source' and not result['warnings']
    section = result['sections'][0]
    assert section['title'] == 'Planter Box Membrane'
    purposes = {p['purpose'] for p in section['parts']}
    assert {'reasoning', 'remediation', 'observation'} <= purposes
    options = [p for p in section['parts'] if p['option_id']]
    assert len(options) >= 2 and len({p['option_id'] for p in options}) == len(options)
    assert len(section['parts']) > 1
    order = [b['sequence'] for p in section['parts'] for b in p['blocks']]
    assert order == sorted(order)
    for part in section['parts']:
        for block in part['blocks']:
            original = source(store, block['anchor_id'])
            assert block['text'] == original['quote'] == original['text']
            assert block['content_kind'] != 'ocr_text'
            assert original['document_id'] == section['document_id']
    assert any(b['media'] for p in section['parts'] for b in p['blocks'])
    assert store.require('SELECT evidence,answer FROM runs WHERE id=?', (rid,)) == before


def test_sections_do_not_read_later_unpublished_units_or_foreign_documents(store, ontology):
    rid = saved_run(store, ontology)
    result = original_report_sections(store, rid)
    section = result['sections'][0]
    unit = store.require('SELECT * FROM units WHERE id=?', (section['parts'][0]['id'],))
    parent_id = 'parent-unpublished'
    block_id = section['parts'][0]['blocks'][0]['id']
    store.insert('parents', {'id': parent_id, 'document_id': unit['document_id'],
                            'text': 'UNPUBLISHED REPLACEMENT', 'block_ids': [block_id], 'dependencies': []})
    store.insert('units', {**unit, 'id': 'unit-unpublished', 'parent_id': parent_id})
    changed = original_report_sections(store, rid)
    assert changed == result
    raw = json.loads(store.require('SELECT evidence FROM runs WHERE id=?', (rid,))['evidence'])
    raw['items'].append({**raw['items'][0], 'id': 'chunkrev-outside-pinned-manifest',
                         'text': 'FOREIGN CLAIM', 'document_id': 'foreign-document'})
    store.execute('UPDATE runs SET evidence=? WHERE id=?', (canonical(raw), rid))
    assert original_report_sections(store, rid) == result


def test_report_section_endpoint_exposes_originals_without_a_model_call(store, ontology):
    rid = saved_run(store, ontology)
    with TestClient(create_app(store.config)) as client:
        response = client.get('/api/runs/'+rid+'/report-sections')
        assert response.status_code == 200 and response.json()['sections']
        assert response.json()['manifest_id'] == store.active_manifest()
        assert client.get('/api/runs/nonexistent/report-sections').status_code == 404
    assert not store.all("SELECT sequence FROM events WHERE owner_id=? AND type='model_started'", (rid,))


def test_empty_run_has_no_invented_sections(store):
    rid = run_record(store)
    assert original_report_sections(store, rid)['sections'] == []


def test_reference_context_returns_retrieved_paragraphs_not_the_entire_book(store, ontology):
    document = Document()
    document.add_heading('Original reference source', 1)
    for index in range(50):
        document.add_paragraph(f'Clause {index}: ' + 'Waterstop threshold requirements retain their stated qualifications. ' * 4)
    payload = io.BytesIO()
    document.save(payload)
    imported = register(payload.getvalue(), 'reference.docx', 'original_reference', store)
    doc = store.require('SELECT * FROM documents WHERE id=?', (imported['document_id'],))
    persist_ir(store, parse_docx(store, doc))
    build_units(store, doc['id'], ontology)
    annotate_domain(store, doc['id'], ontology)
    build_chunks(store, doc['id'], ontology.version)
    projection = stage(store, doc['id'], ontology)
    store.activate(projection['projection_id'], {'profile': projection['profile']})
    hit = search(store, ontology, 'Clause 24 waterstop', limit=1)['items'][0]
    rid = run_record(store)
    store.execute('UPDATE runs SET evidence=? WHERE id=?', (canonical({'items': [hit]}), rid))
    result = original_report_sections(store, rid)
    section = result['sections'][0]
    blocks = [b for part in section['parts'] for b in part['blocks']]
    allowed = {s['block_id'] for s in hit['metadata']['source_segments']}
    assert section['scope'] == 'retrieved_source_passages' and not section['heading_anchor_id']
    assert {b['id'] for b in blocks} <= allowed and len(blocks) < 50
    # Current document metadata cannot relabel the pinned publication's role.
    store.execute("UPDATE documents SET role='company_report' WHERE id=?", (doc['id'],))
    assert original_report_sections(store, rid) == result
