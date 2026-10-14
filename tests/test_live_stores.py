import pytest
import io
from docx import Document
from app.settings import settings
from app.projections import NeoProjection,QdrantProjection
from app.retrieval import search
from tests.conftest import publish_fixture,synthetic_docx


@pytest.mark.skipif(not settings.neo4j_uri or not settings.qdrant_url,reason='Configured Neo4j/Qdrant services required')
def test_actual_neo4j_qdrant_projection_receipts_and_manifest_scope(store,ontology):
    for key in ['neo4j_uri','neo4j_user','neo4j_password','qdrant_url','qdrant_key','embedding_path','tokenizer_path']:
        setattr(store.config,key,getattr(settings,key))
    doc=Document(io.BytesIO(synthetic_docx()))
    doc.add_heading('Separate roof issue',2)
    doc.add_paragraph('Sample Findings:','CHeading')
    doc.add_paragraph('DS-900, separate roof')
    doc.add_paragraph('Roof damage belongs to a separate assessment item.')
    doc.add_paragraph('Reasoning:','CHeading')
    doc.add_paragraph('This roof damage is unrelated to the planter membrane.')
    raw=io.BytesIO()
    doc.save(raw)
    r=publish_fixture(store,ontology,raw.getvalue())
    projection=store.require('SELECT * FROM projections WHERE document_id=? ORDER BY rowid DESC LIMIT 1',(r['document_id'],))
    assert projection['profile']=='full'
    item=store.require("SELECT id FROM items WHERE document_id=? AND title='Planter Box Membrane'",(r['document_id'],))['id']
    paths=NeoProjection(store).neighbors([item],[projection['id']])
    assert paths and all(p['assertion_id'] for p in paths)
    graph=NeoProjection(store).explore([item],[projection['id']],50)
    assert {'Assertion','SourceAnchor'} <= {n['kind'] for n in graph['nodes']}
    assert len(graph['nodes'])<=50
    graph_ids={n['id'] for n in graph['nodes']}
    assert all(e['source'] in graph_ids and e['target'] in graph_ids for e in graph['edges'])
    unrelated=store.require("SELECT id FROM items WHERE document_id=? AND title='Separate roof issue'",(r['document_id'],))['id']
    assert unrelated not in graph_ids
    from app.domain import published_snapshots
    snapshot=published_snapshots(store)[0]
    option=next(r for r in snapshot['records'] if r['class_id']=='RemediationOption')
    typed=NeoProjection(store).explore([option['id']],[projection['id']],50)
    assert {'OntologyClass','DomainAssertion'} <= {n['kind'] for n in typed['nodes']}
    assert any(n['kind'] in {'SourceAnchor','TextPassage'} for n in typed['nodes'])
    assert any(n.get('domain',{}).get('class_id')=='RemediationOption' for n in typed['nodes'] if n.get('domain'))
    assert all(n['assertion']['anchor_ids'] or n['assertion']['block_ids'] for n in typed['nodes'] if n['kind']=='DomainAssertion')
    vectors=QdrantProjection(store).search('waterproofing details for a planter',[projection['id']])
    assert vectors and all(store.one('SELECT id FROM chunk_revisions WHERE id=?',(v['id'],)) for v in vectors)
    result=search(store,ontology,'waterproofing details for a planter')
    assert result['items'] and any('semantic' in e['retrieval_methods'] for e in result['items'])
